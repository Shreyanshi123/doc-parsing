"""Generate small test documents in ./samples (digital PDF, scanned PDF, mixed PDF, docx, pptx, xlsx, csv, png).

    python scripts/make_samples.py
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "samples"

TITLE = "Quarterly Sales Report"
PARA = ("This report summarises the sales performance of the company in the third quarter. "
        "Revenue grew steadily across all regions and customer satisfaction improved.")
TABLE = [["Region", "Q1", "Q2", "Q3"], ["North", "120", "135", "150"], ["South", "98", "110", "121"], ["West", "75", "80", "95"]]


def _font(size):
    from PIL import ImageFont

    for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/Library/Fonts/Arial.ttf", "C:/Windows/Fonts/arial.ttf"):
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def page_image(w=1240, h=1754):
    """A page rendered as a picture (what a scanner produces)."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((100, 120), TITLE, fill="black", font=_font(64))
    y = 260
    words, line = PARA.split(), ""
    for wd in words:
        if len(line + wd) > 55:
            d.text((100, y), line, fill="black", font=_font(34)); y += 52; line = ""
        line += wd + " "
    d.text((100, y), line, fill="black", font=_font(34))
    return img


def make_pdfs():
    # scanned PDF (image only) - needs only Pillow
    page_image().save(OUT / "scanned.pdf", "PDF", resolution=150)
    print("scanned.pdf")
    try:
        import fitz
    except ImportError:
        print("PyMuPDF not installed -> skipping digital.pdf / mixed.pdf (pip install pymupdf)")
        return
    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((72, 90), TITLE, fontsize=24)
    p.insert_textbox(fitz.Rect(72, 120, 520, 220), PARA, fontsize=11)
    y = 260
    for r in TABLE:                       # simple ruled table so find_tables() can detect it
        for i, c in enumerate(r):
            rect = fitz.Rect(72 + i * 100, y, 172 + i * 100, y + 24)
            p.draw_rect(rect, color=(0, 0, 0), width=0.6)
            p.insert_text((rect.x0 + 6, rect.y0 + 16), c, fontsize=10)
        y += 24
    doc.save(OUT / "digital.pdf")
    print("digital.pdf")
    # mixed = digital page + scanned page
    img = io.BytesIO(); page_image().save(img, "PNG")
    p2 = doc.new_page()
    p2.insert_image(p2.rect, stream=img.getvalue())
    doc.save(OUT / "mixed.pdf")
    print("mixed.pdf")


def make_office():
    import docx
    d = docx.Document()
    d.add_heading(TITLE, 0)
    d.add_heading("Summary", 1)
    d.add_paragraph(PARA)
    d.add_paragraph("North region led growth", style="List Bullet")
    d.add_paragraph("West region improved margins", style="List Bullet")
    t = d.add_table(rows=len(TABLE), cols=4)
    for i, r in enumerate(TABLE):
        for j, c in enumerate(r):
            t.cell(i, j).text = c
    d.save(OUT / "report.docx"); print("report.docx")

    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = TITLE
    s.placeholders[1].text = "Revenue up 12%\nCustomer satisfaction improved"
    s.notes_slide.notes_text_frame.text = "Mention the West region."
    prs.save(OUT / "deck.pptx"); print("deck.pptx")

    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sales"
    for r in TABLE:
        ws.append(r)
    wb.create_sheet("Notes").append(["Comment", "Prepared by finance"])
    wb.save(OUT / "sales.xlsx"); print("sales.xlsx")

    (OUT / "sales.csv").write_text("\n".join(",".join(r) for r in TABLE), encoding="utf-8"); print("sales.csv")
    page_image().save(OUT / "scan.png"); print("scan.png")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_pdfs()
    make_office()
    print(f"\nSamples written to {OUT}")

"""XLSX (openpyxl) and CSV/TSV. Each sheet becomes a page holding one Markdown table."""
from __future__ import annotations

import csv
import io

from app.models import TABLE, TEXT, Element, Page, ParsedDocument
from app.parsers.base import BaseParser
from app.parsers.registry import register
from app.parsers.utils import limit_pages, rows_to_markdown

MAX_ROWS = 2000  # POC safety cap per sheet


@register
class SpreadsheetParser(BaseParser):
    name = "spreadsheet"
    description = "openpyxl for XLSX (one page per sheet), csv module for CSV/TSV. Rows capped at 2000/sheet."
    supported_types = {"xlsx", "csv"}

    @classmethod
    def is_available(cls):
        try:
            import openpyxl  # noqa: F401
            return True, ""
        except ImportError:
            return False, "pip install openpyxl"

    def parse(self, data, file_name, file_type, options):
        out = ParsedDocument(file_name, file_type, self.name)
        if file_type == "csv":
            text = self._decode(data)
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(io.StringIO(text), dialect))
            self._add_sheet(out, 1, "", rows)
            return out

        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        names = wb.sheetnames
        for i in range(limit_pages(len(names), options.max_pages)):
            ws = wb[names[i]]
            rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 1)]
            self._add_sheet(out, i + 1, names[i], rows)
        out.warnings.append("Formulas are read as their last cached values (data_only=True).")
        return out

    def _add_sheet(self, out, n, title, rows):
        els = []
        if title:
            els.append(Element(TEXT, f"## {title}"))
        if len(rows) > MAX_ROWS:
            rows = rows[:MAX_ROWS]
            out.warnings.append(f"Sheet '{title or n}' truncated to {MAX_ROWS} rows.")
        md = rows_to_markdown(rows)
        if md:
            els.append(Element(TABLE, md, metadata={"sheet": title}))
        out.pages.append(Page(n, els))

    @staticmethod
    def _decode(data: bytes) -> str:
        for enc in ("utf-8-sig", "utf-16", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeError:
                continue
        return data.decode("utf-8", errors="replace")


@register
class TextParser(BaseParser):
    name = "text"
    description = "Plain text / Markdown passthrough."
    supported_types = {"text"}

    def parse(self, data, file_name, file_type, options):
        out = ParsedDocument(file_name, file_type, self.name)
        out.pages.append(Page(1, [Element(TEXT, SpreadsheetParser._decode(data))]))
        return out

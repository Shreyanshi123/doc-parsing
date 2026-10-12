"""python-pptx: one slide = one page. Title, bullets (with indentation), tables, pictures, notes."""
from __future__ import annotations

from app.models import IMAGE, TABLE, TEXT, Element, Page, ParsedDocument
from app.parsers.base import BaseParser
from app.parsers.registry import register
from app.parsers.utils import limit_pages, rows_to_markdown


@register
class PptxParser(BaseParser):
    name = "pptx"
    description = "python-pptx native parser: one page per slide, titles, bullets, tables, notes, pictures."
    supported_types = {"pptx"}

    @classmethod
    def is_available(cls):
        try:
            import pptx  # noqa: F401
            return True, ""
        except ImportError:
            return False, "pip install python-pptx"

    def parse(self, data, file_name, file_type, options):
        import io

        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        out = ParsedDocument(file_name, file_type, self.name)
        slides = list(prs.slides)
        for i in range(limit_pages(len(slides), options.max_pages)):
            slide = slides[i]
            els: list[Element] = []
            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.has_text_frame and title_shape.text_frame.text.strip():
                els.append(Element(TEXT, f"## {title_shape.text_frame.text.strip()}"))
            shapes = sorted(self._flatten(slide.shapes), key=lambda s: (s.top or 0, s.left or 0))
            for sh in shapes:
                if title_shape is not None and sh.shape_id == title_shape.shape_id:
                    continue
                if sh.has_text_frame:
                    lines = []
                    for para in sh.text_frame.paragraphs:
                        t = "".join(r.text for r in para.runs).strip()
                        if t:
                            lines.append(f"{'  ' * para.level}- {t}")
                    if lines:
                        els.append(Element(TEXT, "\n".join(lines)))
                elif getattr(sh, "has_table", False) and sh.has_table:
                    rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                    md = rows_to_markdown(rows)
                    if md:
                        els.append(Element(TABLE, md))
                elif options.extract_images and sh.shape_type == 13:  # PICTURE
                    try:
                        els.append(Element(IMAGE, "", sh.image.blob, sh.image.ext))
                    except Exception:
                        pass
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    els.append(Element(TEXT, f"> **Speaker notes:** {notes}", metadata={"notes": True}))
            out.pages.append(Page(i + 1, els))
        return out

    def _flatten(self, shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # GROUP
                yield from self._flatten(sh.shapes)
            else:
                yield sh

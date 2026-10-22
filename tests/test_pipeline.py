"""Run with:  python -m unittest discover -s tests -v     (or: pytest)

Tests that need optional pieces (PyMuPDF, tesseract) skip themselves when missing."""
from __future__ import annotations

import importlib.util
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("DATA_DIR", tempfile.mkdtemp(prefix="docpoc_"))

from app.config import settings  # noqa: E402
from app.jobs.store import COMPLETED, FAILED, JobStore  # noqa: E402
from app.jobs.worker import process_job  # noqa: E402
from app.models import ParseOptions  # noqa: E402
from app.normalizer import normalize_markdown  # noqa: E402
from app.parsers import list_parsers  # noqa: E402
from app.pipeline import ParseFailed, build_output, parse_bytes  # noqa: E402
from app.routing.analyzer import DocProfile  # noqa: E402
from app.routing.file_type import UnsupportedFileType, resolve_file_type  # noqa: E402
from app.routing.router import NoParserAvailable, route  # noqa: E402
from app.storage.base import StorageError  # noqa: E402
from app.storage.local import LocalStorage  # noqa: E402
from app.storage.output_store import get_output_store  # noqa: E402

HAS_FITZ = importlib.util.find_spec("fitz") is not None
HAS_TESS = shutil.which("tesseract") is not None


def _docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Title", 0)
    d.add_heading("Section", 1)
    d.add_paragraph("Hello world")
    t = d.add_table(rows=2, cols=2)
    for i, row in enumerate([["a", "b"], ["1", "2"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S1"
    ws.append(["x", "y"]); ws.append([1, 2])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Deck title"
    s.placeholders[1].text = "bullet one"
    buf = io.BytesIO(); prs.save(buf)
    return buf.getvalue()


class FileTypeTests(unittest.TestCase):
    def test_magic_bytes_beat_extension(self):
        self.assertEqual(resolve_file_type(_docx_bytes(), "wrong.pdf"), "docx")
        self.assertEqual(resolve_file_type(b"%PDF-1.7 ...", "x.bin"), "pdf")
        self.assertEqual(resolve_file_type(_xlsx_bytes(), "a.xlsx"), "xlsx")
        self.assertEqual(resolve_file_type(_pptx_bytes(), "a.pptx"), "pptx")

    def test_unsupported(self):
        with self.assertRaises(UnsupportedFileType):
            resolve_file_type(b"random bytes", "file.xyz")
        with self.assertRaises(UnsupportedFileType):
            resolve_file_type(b"\xd0\xcf\x11\xe0 legacy", "old.doc")


class NormalizerTests(unittest.TestCase):
    def test_cleanup(self):
        out = normalize_markdown("a  b \r\n\r\n\r\n\r\nc\u00a0d\x00\nText\n# Head")
        self.assertNotIn("\r", out)
        self.assertNotIn("\n\n\n", out)
        self.assertIn("a b", out)
        self.assertIn("Text\n\n# Head", out)


class ParserTests(unittest.TestCase):
    def test_docx(self):
        r = parse_bytes(_docx_bytes(), "x.docx")
        self.assertEqual(r.routing.selected, "docx")
        md = r.document.markdown
        self.assertIn("# Title", md); self.assertIn("# Section", md)
        self.assertIn("| a | b |", md)
        self.assertEqual(r.document.stats()["tables"], 1)

    def test_xlsx_and_csv(self):
        r = parse_bytes(_xlsx_bytes(), "x.xlsx")
        self.assertIn("## S1", r.document.markdown); self.assertIn("| 1 | 2 |", r.document.markdown)
        r = parse_bytes(b"h1,h2\n1,2\n", "x.csv")
        self.assertIn("| h1 | h2 |", r.document.markdown)

    def test_pptx(self):
        r = parse_bytes(_pptx_bytes(), "x.pptx")
        self.assertIn("## Deck title", r.document.markdown)
        self.assertIn("bullet one", r.document.markdown)

    def test_text(self):
        r = parse_bytes("héllo".encode(), "a.txt")
        self.assertEqual(r.routing.selected, "text")

    @unittest.skipUnless(HAS_TESS, "tesseract not installed")
    def test_image_ocr(self):
        from PIL import Image, ImageDraw, ImageFont

        img = Image.new("RGB", (900, 200), "white")
        font = None
        for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/Library/Fonts/Arial.ttf", "C:/Windows/Fonts/arial.ttf"):
            if Path(p).exists():
                font = ImageFont.truetype(p, 60); break
        ImageDraw.Draw(img).text((30, 50), "Invoice 12345", fill="black", font=font)
        buf = io.BytesIO(); img.save(buf, "PNG")
        r = parse_bytes(buf.getvalue(), "i.png")
        self.assertEqual(r.routing.selected, "ocr")
        self.assertIn("12345", r.document.markdown)

    def test_forced_parser_errors(self):
        with self.assertRaises(NoParserAvailable):
            parse_bytes(_docx_bytes(), "x.docx", parser="pymupdf")   # does not support docx
        with self.assertRaises(KeyError):
            parse_bytes(_docx_bytes(), "x.docx", parser="nope")

    def test_empty_doc_fails_cleanly(self):
        with self.assertRaises(ParseFailed):
            parse_bytes(b"   \n  ", "empty.txt")

    def test_output_contract(self):
        out = build_output(parse_bytes(_docx_bytes(), "x.docx"), file_id="f1", tenant_id="t1")
        for key in ("metadata", "content", "pages", "routing", "stats", "timings_ms", "warnings"):
            self.assertIn(key, out)
        self.assertEqual(out["metadata"]["tenant_id"], "t1")
        el = out["pages"][0]["elements"][0]
        self.assertIn(el["content_type"], ("text", "table", "image"))
        self.assertEqual(out["pages"][0]["page_number"], 1)
        json.dumps(out)  # must be serialisable


class RegistryTests(unittest.TestCase):
    EXPECTED = {"pymupdf", "pymupdf4llm", "pdfminer", "pdfplumber", "ocr", "easyocr", "hybrid_pdf", "vlm",
                "opendataloader", "docling", "unstructured", "markitdown", "llamaparse", "azure_di",
                "docx", "pptx", "spreadsheet", "text", "smart_pdf",
                "langchain_loader", "llamaindex_reader", "haystack_converter"}

    def test_all_parsers_registered_and_report_availability(self):
        infos = {p["name"]: p for p in list_parsers()}
        self.assertEqual(self.EXPECTED, set(infos))
        for p in infos.values():  # unavailable parsers must explain why, never raise
            self.assertTrue(p["available"] or p["unavailable_reason"])

    @unittest.skipUnless(importlib.util.find_spec("pdfplumber") and importlib.util.find_spec("reportlab"),
                         "needs pdfplumber + reportlab (test-only PDF generator)")
    def test_pdfplumber_forced(self):
        r = parse_bytes(_simple_pdf(), "t.pdf", parser="pdfplumber")
        self.assertIn("Hello PDF", r.document.markdown)

    def test_unstructured_html_table_helper(self):
        from app.parsers.unstructured_parser import html_table_to_markdown

        md = html_table_to_markdown("<table><tr><th>a</th><th>b</th></tr><tr><td>1</td><td>2</td></tr></table>")
        self.assertIn("| a | b |", md); self.assertIn("| 1 | 2 |", md)


def _simple_pdf() -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO(); c = canvas.Canvas(buf); c.drawString(72, 750, "Hello PDF world"); c.save()
    return buf.getvalue()


class RouterTests(unittest.TestCase):
    def _p(self, ft, cls):
        return DocProfile(file_type=ft, size_bytes=1, classification=cls)

    def test_pdf_rules(self):
        names = {p["name"]: p["available"] for p in list_parsers()}
        if names.get("pymupdf"):
            self.assertEqual(route(self._p("pdf", "digital"), ParseOptions()).selected, "pymupdf")
        if names.get("pymupdf4llm") and HAS_FITZ:   # scanned: pymupdf4llm scored 93%, plain OCR 49% in our test
            self.assertEqual(route(self._p("pdf", "scanned"), ParseOptions()).selected, "pymupdf4llm")
        if names.get("hybrid_pdf"):
            self.assertEqual(route(self._p("pdf", "mixed"), ParseOptions()).selected, "hybrid_pdf")

    def test_pdf_with_tables_prefers_pymupdf4llm(self):
        names = {p["name"]: p["available"] for p in list_parsers()}
        if names.get("pymupdf4llm"):
            prof = DocProfile(file_type="pdf", size_bytes=1, classification="digital", traits=["tables"], est_tables=3)
            self.assertEqual(route(prof, ParseOptions()).selected, "pymupdf4llm")

    def test_prefer_quality_uses_docling_for_digital_and_scanned(self):
        names = {p["name"]: p["available"] for p in list_parsers()}
        if names.get("docling"):
            for cls in ("digital", "scanned"):
                self.assertEqual(route(self._p("pdf", cls), ParseOptions(prefer_quality=True)).selected, "docling")

    def test_figures_add_a_note_but_do_not_change_the_pick(self):
        prof = DocProfile(file_type="pdf", size_bytes=1, classification="digital", traits=["figures"],
                          figure_count=2, figure_pages=1)
        d = route(prof, ParseOptions())
        self.assertEqual(d.selected, "pymupdf")
        self.assertTrue(any("does not read chart values" in n for n in d.notes))

    def test_office_rules(self):
        self.assertEqual(route(self._p("docx", "native_office"), ParseOptions()).selected, "docx")
        self.assertEqual(route(self._p("xlsx", "native_office"), ParseOptions()).selected, "spreadsheet")

    def test_override(self):
        d = route(self._p("docx", "native_office"), ParseOptions(), "docx")
        self.assertTrue(d.overridden)


class StorageTests(unittest.TestCase):
    def test_local_allowlist(self):
        with self.assertRaises(StorageError):
            LocalStorage().read("/etc/hosts")
        p = settings.uploads_dir / "ok.txt"
        p.write_text("hi")
        self.assertEqual(LocalStorage().read(str(p)), b"hi")


class JobFlowTests(unittest.TestCase):
    def test_job_completes_and_fails(self):
        store = JobStore(Path(tempfile.mkdtemp()) / "jobs.db")
        src = settings.uploads_dir / "job.docx"
        src.write_bytes(_docx_bytes())
        jid = store.create(tenant_id="t", file_id="f", file_name="job.docx", source_uri=str(src))
        job = store.claim_next()
        self.assertEqual(job["job_id"], jid)
        self.assertIsNone(store.claim_next())          # nothing else queued
        process_job(store, job)
        done = store.get(jid)
        self.assertEqual(done["status"], COMPLETED)
        self.assertEqual(done["engine"], "docx")
        out = json.loads(get_output_store().load(done["output_uri"]))
        self.assertEqual(out["metadata"]["file_id"], "f")

        bad = store.create(tenant_id="t", file_id="g", file_name="x.docx", source_uri="/etc/passwd")
        process_job(store, store.claim_next())
        self.assertEqual(store.get(bad)["status"], FAILED)


@unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
class PdfTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        sys.path.insert(0, str(ROOT / "scripts"))
        import make_samples

        make_samples.OUT = Path(tempfile.mkdtemp())
        make_samples.make_pdfs()
        cls.dir = make_samples.OUT

    def test_digital(self):
        r = parse_bytes((self.dir / "digital.pdf").read_bytes(), "digital.pdf")
        self.assertIn("tables", r.profile.traits)                    # the sample has a real ruled table
        self.assertIn(r.routing.selected, ("pymupdf4llm", "pymupdf"))
        self.assertIn("Quarterly Sales Report", r.document.markdown)
        self.assertIn("North", r.document.markdown)                  # the table content is there

    @unittest.skipUnless(HAS_TESS, "tesseract not installed")
    def test_scanned_and_mixed(self):
        r = parse_bytes((self.dir / "scanned.pdf").read_bytes(), "scanned.pdf")
        self.assertIn(r.routing.selected, ("pymupdf4llm", "ocr"))     # ocr if pymupdf4llm returns nothing
        self.assertIn("Quarterly", r.document.markdown)
        r = parse_bytes((self.dir / "mixed.pdf").read_bytes(), "mixed.pdf")
        self.assertEqual(r.routing.selected, "hybrid_pdf")
        self.assertEqual(r.document.metadata["ocr_pages"], [2])


if __name__ == "__main__":
    unittest.main()
